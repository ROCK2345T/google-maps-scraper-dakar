"""
Module d'exportation des données scrappées vers Excel (.xlsx) et CSV.
Garantit une mise en page soignée, le support de l'encodage UTF-8 (accents)
et le dimensionnement automatique des colonnes sous Microsoft Excel.
"""

import io
import os
import re
from typing import List, Dict, Any, Union
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Définition des colonnes standard pour l'export
COLUMN_MAPPING = {
    "name": "Nom de l'entreprise",
    "category": "Catégorie",
    "phone": "Téléphone",
    "primary_email": "Email de contact",
    "all_emails": "Tous les emails",
    "website": "Site Web",
    "rating": "Note (/5)",
    "reviews_count": "Nombre d'avis",
    "address": "Adresse complète",
    "latitude": "Latitude",
    "longitude": "Longitude",
    "instagram": "Instagram",
    "linkedin": "LinkedIn",
    "facebook": "Facebook",
    "twitter": "Twitter / X",
    "maps_url": "Lien Google Maps"
}

def format_records_for_export(records: List[Dict[str, Any]]) -> pd.DataFrame:
    """Transforme une liste de dictionnaires en DataFrame Pandas avec colonnes bien nommées."""
    if not records:
        return pd.DataFrame(columns=list(COLUMN_MAPPING.values()))
        
    formatted = []
    for r in records:
        # Formater les emails
        emails_list = r.get("emails", [])
        if isinstance(emails_list, list):
            emails_str = ", ".join(emails_list)
        else:
            emails_str = str(emails_list or "")
            
        row = {
            COLUMN_MAPPING["name"]: r.get("name", ""),
            COLUMN_MAPPING["category"]: r.get("category", ""),
            COLUMN_MAPPING["phone"]: r.get("phone", ""),
            COLUMN_MAPPING["primary_email"]: r.get("primary_email", "") or (emails_list[0] if emails_list else ""),
            COLUMN_MAPPING["all_emails"]: emails_str,
            COLUMN_MAPPING["website"]: r.get("website", ""),
            COLUMN_MAPPING["rating"]: r.get("rating", ""),
            COLUMN_MAPPING["reviews_count"]: r.get("reviews_count", ""),
            COLUMN_MAPPING["address"]: r.get("address", ""),
            COLUMN_MAPPING["latitude"]: r.get("latitude", ""),
            COLUMN_MAPPING["longitude"]: r.get("longitude", ""),
            COLUMN_MAPPING["instagram"]: r.get("instagram", ""),
            COLUMN_MAPPING["linkedin"]: r.get("linkedin", ""),
            COLUMN_MAPPING["facebook"]: r.get("facebook", ""),
            COLUMN_MAPPING["twitter"]: r.get("twitter", ""),
            COLUMN_MAPPING["maps_url"]: r.get("maps_url", "")
        }
        formatted.append(row)
        
    df = pd.DataFrame(formatted)
    return df

def generate_excel_bytes(records_or_df: Union[List[Dict[str, Any]], pd.DataFrame]) -> bytes:
    """Génère un flux d'octets Excel (.xlsx) stylisé prêt pour le téléchargement."""
    if isinstance(records_or_df, pd.DataFrame):
        df = records_or_df
    else:
        df = format_records_for_export(records_or_df)
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entreprises Scrappées"
    
    # Activer l'affichage des lignes de grille
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Bleu foncé Pro
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10)
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Gris très clair
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # En-têtes
    headers = list(df.columns)
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28
    
    # Lignes de données
    for row_idx, row_data in enumerate(df.values, start=2):
        ws.append(list(row_data))
        use_zebra = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row_idx].height = 22
        
    # Auto-ajustement des largeurs de colonnes
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        # Limiter entre 12 et 50 caractères pour un rendu optimal
        adjusted_width = max(min(max_len + 3, 50), 12)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    # Activer le filtre automatique
    if len(headers) > 0 and len(df) > 0:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(df) + 1}"
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_csv_bytes(records_or_df: Union[List[Dict[str, Any]], pd.DataFrame], sep: str = ";") -> bytes:
    """
    Génère un flux d'octets CSV avec UTF-8 BOM (utf-8-sig) pour ouverture
    directe sans bug d'accentuation dans Microsoft Excel.
    """
    if isinstance(records_or_df, pd.DataFrame):
        df = records_or_df
    else:
        df = format_records_for_export(records_or_df)
        
    csv_str = df.to_csv(index=False, sep=sep, encoding="utf-8-sig")
    return csv_str.encode("utf-8-sig")

def save_export_files(records: List[Dict[str, Any]], base_filename: str, output_dir: str = "exports") -> Dict[str, str]:
    """Sauvegarde les fichiers Excel et CSV dans le dossier spécifié."""
    os.makedirs(output_dir, exist_ok=True)
    clean_name = re.sub(r'[^\w\-_\.]', '_', base_filename)
    
    excel_path = os.path.join(output_dir, f"{clean_name}.xlsx")
    csv_path = os.path.join(output_dir, f"{clean_name}.csv")
    
    excel_bytes = generate_excel_bytes(records)
    with open(excel_path, "wb") as f:
        f.write(excel_bytes)
        
    csv_bytes = generate_csv_bytes(records)
    with open(csv_path, "wb") as f:
        f.write(csv_bytes)
        
    return {
        "excel": excel_path,
        "csv": csv_path
    }
